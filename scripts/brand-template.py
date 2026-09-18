#!/usr/bin/env python3
"""Approved brand tokens -> native POTX with six named placeholder layouts.
No original master is imported. Source company material remains private.
"""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
import argparse,json,sys,zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from brand_common import validate_brand,local_logo

def element(tag,**attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(k,str(v))
    return e

def placeholder(idx,name,kind,box,brand):
    heading=kind in ('title','ctrTitle');type_role='heading' if heading else 'body';t=brand['typography'][type_role]
    sp=element('p:sp');nv=element('p:nvSpPr');nv.append(element('p:cNvPr',id=idx+2,name=name));nv.append(element('p:cNvSpPr'));nvp=element('p:nvPr');nvp.append(element('p:ph',type=kind,idx=idx));nv.append(nvp);sp.append(nv)
    pr=element('p:spPr');xf=element('a:xfrm');xf.append(element('a:off',x=round(box[0]*9525),y=round(box[1]*9525)));xf.append(element('a:ext',cx=round(box[2]*9525),cy=round(box[3]*9525)));pr.append(xf);sp.append(pr)
    tx=element('p:txBody');tx.append(element('a:bodyPr',wrap='square',lIns=0,rIns=0,tIns=0,bIns=0));tx.append(element('a:lstStyle'));para=element('a:p');pp=element('a:pPr',algn='l');pp.append(element('a:buNone'));default=element('a:defRPr',sz=round(t['size']*.75*100));fill=element('a:solidFill');fill.append(element('a:srgbClr',val=brand['colors']['text'][1:]));default.append(fill);default.append(element('a:latin',typeface=t['family'][0]));pp.append(default);para.append(pp);tx.append(para);sp.append(tx);return sp

def create(brand,out,base,content=None,allow_draft=False):
    validate_brand(brand,allow_draft);out=Path(out)
    if out.exists():raise ValueError('Output exists; use a new path')
    if out.suffix not in ('.potx','.pptx'):raise ValueError('Output must be .potx or .pptx')
    logo=local_logo(brand,base);prs=Presentation();w=brand['canvas']['width'];h=brand['canvas']['height'];prs.slide_width=round(w*9525);prs.slide_height=round(h*9525)
    # Start from a valid package; replace only the shapes and styles we explicitly own.
    master=prs.slide_masters[0]
    for shape in list(master.shapes):master.shapes._spTree.remove(shape._element)
    # Reset inherited default Office centering/bullets for future new slides as well.
    # Layout-local text settings alone do not override every client's master defaults.
    txstyles=master._element.find(qn('p:txStyles'))
    if txstyles is not None:master._element.remove(txstyles)
    txstyles=element('p:txStyles')
    for tag,role in [('titleStyle','heading'),('bodyStyle','body'),('otherStyle','body')]:
        style=element('p:'+tag);t=brand['typography'][role]
        for level in range(1,10):
            pp=element('a:lvl'+str(level)+'pPr',algn='l',marL=0,indent=0);pp.append(element('a:buNone'))
            run=element('a:defRPr',sz=round(t['size']*.75*100));run.append(element('a:latin',typeface=t['family'][0]));fill=element('a:solidFill');fill.append(element('a:srgbClr',val=brand['colors']['text'][1:]));run.append(fill);pp.append(run);style.append(pp)
        txstyles.append(style)
    master._element.append(txstyles)
    while len(prs.slide_layouts)>6:prs.slide_layouts.remove(prs.slide_layouts[-1])
    m=brand['spacing']['margin'];g=brand['spacing']['gutter'];usable=w-2*m;body_y=max(m+brand['typography']['heading']['size']*2.25, h*.30);body_h=h-body_y-m
    titles=['Title / Subtitle','Evidence / Explanation','Paired Comparison','Image / Caption','Dense Information','Quiet Close']
    layouts=[
      [('title',(m,h*.24,usable,h*.25)),('subTitle',(m,h*.56,usable,h*.19))],
      [('title',(m,m,usable,h*.17)),('body',(m,body_y,(usable-g)*.68,body_h)),('body',(m+(usable-g)*.68+g,body_y,(usable-g)*.32,body_h))],
      [('title',(m,m,usable,h*.17)),('body',(m,body_y,(usable-g)/2,body_h)),('body',(m+(usable-g)/2+g,body_y,(usable-g)/2,body_h))],
      [('title',(m,m,usable,h*.17)),('pic',(m,body_y,(usable-g)*.66,body_h)),('body',(m+(usable-g)*.66+g,body_y,(usable-g)*.34,body_h))],
      [('title',(m,m,usable,h*.17)),('body',(m,body_y,usable,body_h))],
      [('title',(m,h*.30,usable,h*.29)),('subTitle',(m,h*.66,usable,h*.17))]
    ]
    for layout,name,boxes in zip(prs.slide_layouts,titles,layouts):
        layout.name=name
        for shape in list(layout.shapes):layout.shapes._spTree.remove(shape._element)
        layout.background.fill.solid();layout.background.fill.fore_color.rgb=RGBColor.from_string(brand['colors']['background'][1:])
        for idx,(kind,box) in enumerate(boxes):layout.shapes._spTree.append(placeholder(idx,'Title' if idx==0 else ('Image' if kind=='pic' else f'Content {idx}'),kind,box,brand))
    # Update the theme, so newly inserted native objects inherit approved tokens.
    theme_part=next(p for p in prs.part.package.iter_parts() if str(p.partname)=='/ppt/theme/theme1.xml');theme=etree.fromstring(theme_part.blob)
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'};mapping={'dk1':'text','lt1':'background','dk2':'muted','lt2':'surface','accent1':'accent','accent2':'line'}
    for token,role in mapping.items():
        target=theme.find('.//a:clrScheme/a:'+token,ns)
        if target is not None:
            for c in list(target):target.remove(c)
            target.append(element('a:srgbClr',val=brand['colors'][role][1:]))
    for family,role in [('majorFont','heading'),('minorFont','body')]:
        t=theme.find('.//a:'+family+'/a:latin',ns)
        if t is not None:t.set('typeface',brand['typography'][role]['family'][0])
    theme_part._blob=etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
    rows=content if content is not None else [{'layout':i,'title':t,'body':['Replace with approved audience content.']*(len(layouts[i])-1)} for i,t in enumerate(titles)]
    if not isinstance(rows,list) or not rows:raise ValueError('Content must be a nonempty array')
    for row in rows:
        index=row.get('layout')
        if not isinstance(index,int) or not 0<=index<6:raise ValueError('Layout index must be 0..5')
        if not isinstance(row.get('title'),str) or not row['title'].strip():raise ValueError('A content title is required')
        body=row.get('body',[])
        if not isinstance(body,list) or not all(isinstance(x,str) for x in body):raise ValueError('body must be an array of strings')
        slide=prs.slides.add_slide(prs.slide_layouts[index]);slide.shapes.title.text=row['title']
        pos=0
        for sh in slide.placeholders:
            if sh is slide.shapes.title or sh.placeholder_format.idx==0:continue
            if sh.has_text_frame:
                sh.text=body[pos] if pos<len(body) else '';pos+=1
        # Set explicit run formatting on starter slides as well as layout inheritance.
        for sh in slide.shapes:
            if sh.has_text_frame:
                role='heading' if sh.placeholder_format.idx==0 else 'body'
                for para in sh.text_frame.paragraphs:
                    para.alignment=PP_ALIGN.LEFT
                    pp=para._p.get_or_add_pPr()
                    for child in list(pp):
                        if child.tag in {qn('a:buChar'),qn('a:buAutoNum'),qn('a:buNone')}:pp.remove(child)
                    pp.append(element('a:buNone'))
                    para.space_after=Pt(10);para.line_spacing=brand['typography'][role]['lineHeight']
                    for run in para.runs:run.font.name=brand['typography'][role]['family'][0];run.font.size=Pt(brand['typography'][role]['size']*.75);run.font.color.rgb=RGBColor.from_string(brand['colors']['text'][1:])
        if logo:
            from PIL import Image
            l=brand['logo'];iw,ih=Image.open(logo).size;ratio=min(l['width']/iw,l['height']/ih);lw=iw*ratio;lh=ih*ratio
            slide.shapes.add_picture(str(logo),round((l['x']+(l['width']-lw)/2)*9525),round((l['y']+(l['height']-lh)/2)*9525),width=round(lw*9525),height=round(lh*9525))
    prs.core_properties.title=brand['name']+' template';prs.core_properties.author='';prs.core_properties.last_modified_by='';prs.core_properties.comments='Generated from reviewed identity tokens; no source deck content imported.'
    stream=BytesIO();prs.save(stream);blob=stream.getvalue()
    if out.suffix=='.potx':
        dst=BytesIO()
        with zipfile.ZipFile(BytesIO(blob)) as z,zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as o:
            for name in z.namelist():
                data=z.read(name)
                if name=='[Content_Types].xml':data=data.replace(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',b'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml')
                o.writestr(name,data)
        blob=dst.getvalue()
    out.parent.mkdir(parents=True,exist_ok=True);out.write_bytes(blob)
    return {'layouts':6,'slides':len(prs.slides),'format':out.suffix,'editable':True,'limits':['Six authored layouts, not arbitrary master import.','Native Office rendering must be checked in the intended application.','Image placeholder population requires image authoring; content.json supplies text only.']}
if __name__=='__main__':
    a=argparse.ArgumentParser(description=__doc__);a.add_argument('brand');a.add_argument('--out',required=True);a.add_argument('--content');a.add_argument('--allow-draft',action='store_true');args=a.parse_args()
    try:print(json.dumps(create(json.loads(Path(args.brand).read_text()),args.out,Path(args.brand).parent,json.loads(Path(args.content).read_text()) if args.content else None,args.allow_draft),indent=2))
    except Exception as e:print('Brand template:',e,file=sys.stderr);sys.exit(1)
