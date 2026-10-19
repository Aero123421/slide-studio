#!/usr/bin/env python3
"""Static PDF or picture-based PPTX from authored HTML, without persistent PNGs.
Use export-pptx.mjs for the separate native text/shape scene route.
"""
import argparse,io,json,os,sys,tempfile,unicodedata
from pathlib import Path
from qa import launch,open_deck,guard_network
from font_audit import load_contract,settle_fonts,audit_fonts,print_text,pdf_postflight

# Some fonts (Inter among them) give their alternate glyphs private-use code points: contextual
# punctuation between figures, tabular figures, character variants. Chromium then writes those
# code points into the PDF text layer, so "18:00" cannot be searched or copied. When the strict
# postflight finds that, the PDF is rendered once more with those substitutions turned off;
# layout features (Japanese palt, kerning, ligatures) are kept.
PDF_SAFE_TEXT=r'''()=>{let n=0;const keep=/^["'](palt|halt|vpal|vhal|kern|vkrn|liga|clig|locl|ccmp)["']/;
 for(const el of document.querySelectorAll('*')){const s=getComputedStyle(el),f=s.fontFeatureSettings;
  if(f&&f!=='normal'){const k=f.split(',').map(x=>x.trim()).filter(x=>keep.test(x));el.style.setProperty('font-feature-settings',k.length?k.join(', '):'normal','important')}
  el.style.setProperty('font-variant-ligatures','no-contextual','important');el.style.setProperty('font-variant-numeric','normal','important');n++}
 return n}'''

# CJK Radicals Supplement → the unified ideograph drawn with the same glyph (Unicode
# EquivalentUnifiedIdeograph). Some CJK fonts map both code points to one glyph and Chromium's
# PDF text layer then names the radical (⻑ for 長), so "店長" cannot be searched.
RADICAL_IDEOGRAPH={0x2E8C:0x5C0F,0x2E97:0x5FC3,0x2E9C:0x65E5,0x2E9D:0x6708,0x2E9F:0x6BCD,0x2EA0:0x6C11,0x2EA7:0x725B,
 0x2EA9:0x738B,0x2EAC:0x793A,0x2EAE:0x7AF9,0x2EB7:0x7F8A,0x2EBC:0x8089,0x2EC1:0x864E,0x2EC4:0x897F,0x2EC5:0x89C1,
 0x2EC6:0x89D2,0x2EC9:0x8D1D,0x2ECA:0x8DB3,0x2ECB:0x8F66,0x2ED1:0x9577,0x2ED3:0x957F,0x2ED4:0x95E8,0x2ED7:0x96E8,
 0x2ED8:0x9752,0x2ED9:0x97E6,0x2EDA:0x9875,0x2EDB:0x98CE,0x2EDC:0x98DE,0x2EDD:0x98DF,0x2EE2:0x9A6C,0x2EE3:0x9AA8,
 0x2EE4:0x9B3C,0x2EE5:0x9C7C,0x2EE6:0x9E1F,0x2EE8:0x9EA6,0x2EE9:0x9EC4,0x2EEB:0x6589,0x2EEC:0x9F50,0x2EED:0x6B6F,
 0x2EEE:0x9F7F,0x2EEF:0x7ADC,0x2EF0:0x9F99,0x2EF1:0x9F9C,0x2EF2:0x4E80,0x2EF3:0x9F9F}

def repair_text_layer(path):
    """Rewrite ToUnicode targets that name a CJK or Kangxi radical to its ideograph. Glyphs are untouched."""
    import re
    from pypdf import PdfReader,PdfWriter
    def fix_targets(text):
        def hexsub(h):
            v=int(h,16) if len(h)==4 else None
            if v in RADICAL_IDEOGRAPH:return f'{RADICAL_IDEOGRAPH[v]:04X}'
            if v is not None and 0x2F00<=v<=0x2FDF:  # Kangxi Radicals: NFKC gives the ideograph (⽉ → 月)
                c=unicodedata.normalize('NFKC',chr(v))
                if len(c)==1 and ord(c)<=0xFFFF:return f'{ord(c):04X}'
            return h
        out=[];block=None;n=0
        for line in text.split('\n'):
            t=line.strip()
            if t.endswith('beginbfchar'):block='char'
            elif t.endswith('beginbfrange'):block='range'
            elif t in ('endbfchar','endbfrange'):block=None
            elif block=='char':
                m=re.fullmatch(r'(<[0-9A-Fa-f]+>\s*)<([0-9A-Fa-f]+)>',t)
                if m and hexsub(m.group(2))!=m.group(2):line=line.replace('<'+m.group(2)+'>','<'+hexsub(m.group(2))+'>');n+=1
            elif block=='range':
                m=re.fullmatch(r'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*(.*)',t)
                if m and m.group(1)==m.group(2):  # a one-code range is a plain mapping
                    rest=re.sub(r'<([0-9A-Fa-f]+)>',lambda k:'<'+hexsub(k.group(1))+'>',m.group(3))
                    if rest!=m.group(3):line=line.replace(m.group(3),rest);n+=1
            out.append(line)
        return '\n'.join(out),n
    reader=PdfReader(path);writer=PdfWriter(clone_from=reader);seen=set();total=0
    def fonts_of(res,depth=0):
        if not res or depth>4:return
        res=res.get_object()
        for ref in (res.get('/Font') or {}).values():yield ref.get_object()
        for ref in (res.get('/XObject') or {}).values():
            x=ref.get_object()
            if x.get('/Subtype')=='/Form':yield from fonts_of(x.get('/Resources'),depth+1)
    for page in writer.pages:
        for font in fonts_of(page.get('/Resources')):
            tu=font.get('/ToUnicode')
            if tu is None:continue
            obj=tu.get_object()
            if id(obj) in seen:continue
            seen.add(id(obj));text=obj.get_data().decode('latin-1');fixed,n=fix_targets(text)
            if n:obj.set_data(fixed.encode('latin-1'));total+=n
    if total:
        with open(path,'wb') as f:writer.write(f)
    return total

def run(a):
    from playwright.sync_api import sync_playwright
    inp=Path(a.input).resolve();dest=a.output or a.out
    if not dest:raise ValueError('Specify OUTPUT or --out OUTPUT')
    if a.output and a.out:raise ValueError('Use OUTPUT or --out, not both')
    out=Path(dest).absolute()
    if out.exists() and not a.force:raise ValueError('Output exists. Use --force only to replace the requested file.')
    if out.is_symlink():raise ValueError('Symlink output rejected')
    if out==inp:raise ValueError('Output cannot replace input HTML')
    if out.suffix.lower() not in ['.pdf','.pptx']:raise ValueError('Output must end in .pdf or .pptx')
    contract=load_contract(getattr(a,'font_contract',None));font_report=None;postflight=None
    if getattr(a,'require_font_contract',False) and contract is None:raise ValueError('This delivery requires an approved --font-contract')
    errors=[];blocked=[];out.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix='.studio-export-',suffix=out.suffix,dir=out.parent);os.close(fd)
    try:
        with sync_playwright() as p:
            browser=launch(p)
            try:
                page=browser.new_page(viewport={'width':1280,'height':780},service_workers='block',reduced_motion='reduce');page.on('pageerror',lambda e:errors.append(str(e)));guard_network(page,blocked);open_deck(page,inp)
                meta=page.evaluate("window.slideStudio.meta || {width:1280,height:720,slides:[]}");w=meta.get('width',1280);h=meta.get('height',720)
                page.evaluate('()=>window.slideStudio.prepareExport?.() || window.slideStudio.showAllBuilds()')
                n=page.locator('.slide').count()
                if out.suffix.lower()=='.pdf':
                    page.emulate_media(media='print')
                    settle_fonts(page)
                    font_report=audit_fonts(page,contract)
                    if not font_report['passed']:raise ValueError('Unapproved print fonts: '+json.dumps(font_report['errors'],ensure_ascii=False))
                    expected=print_text(page)
                    page.pdf(path=tmp,print_background=True,prefer_css_page_size=True)
                    repaired=repair_text_layer(tmp)
                    postflight=pdf_postflight(tmp,expected)
                    if not postflight['passed'] and postflight.get('privateUse'):
                        first=postflight['errors'];page.evaluate(PDF_SAFE_TEXT);settle_fonts(page)
                        page.pdf(path=tmp,print_background=True,prefer_css_page_size=True)
                        repaired=repair_text_layer(tmp)
                        postflight=pdf_postflight(tmp,expected)
                        postflight['textLayerFallback']={'reason':first[:3],'applied':'Contextual alternates, tabular figures and character variants were turned off for this PDF so its text can be searched and copied. Look at the PDF: digits and punctuation may be set slightly differently from the HTML.'}
                    if repaired:postflight['radicalsRepaired']=repaired
                    if not postflight['passed']:raise ValueError('PDF postflight failed: '+repr(postflight['errors']))
                else:
                    from pptx import Presentation
                    from pptx.util import Inches
                    prs=Presentation();prs.slide_width=Inches(w/96);prs.slide_height=Inches(h/96);prs.core_properties.author='';prs.core_properties.title=meta.get('title','Presentation')
                    for i in range(n):
                        page.evaluate('(i)=>window.slideStudio.show(i,99999)',i);page.evaluate('()=>new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)))')
                        image=page.locator('.slide').nth(i).screenshot(animations='disabled')
                        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(io.BytesIO(image),0,0,width=prs.slide_width,height=prs.slide_height)
                        notes=meta.get('slides',[]);note=notes[i] if len(notes)>i else {}
                        slide.notes_slide.notes_text_frame.text=note.get('notes','')+'\n\nSources: '+json.dumps(note.get('sources',[]),ensure_ascii=False)+'\nStatic picture-based export. HTML interactivity and motion are not preserved.'
                    prs.save(tmp)
                if errors or blocked:raise ValueError('Export blocked by page errors or external assets: '+repr(errors+blocked))
            finally:browser.close()
        os.replace(tmp,out)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
    return {'output':str(out),'slides':n,'mode':'static PDF' if out.suffix.lower()=='.pdf' else 'fidelity PPTX (pictures; not native editable objects)','persistentPNG':0,'fontAudit':font_report,'pdfPostflight':postflight,'deliveryApproved':False,'limits':'Export/font/text integrity only. Review the rendered PDF and the underlying claims before delivery.'}
p=argparse.ArgumentParser(description=__doc__);p.add_argument('input');p.add_argument('output',nargs='?');p.add_argument('--out');p.add_argument('--force',action='store_true');p.add_argument('--font-contract');p.add_argument('--require-font-contract',action='store_true')
if __name__=='__main__':
    try:print(json.dumps(run(p.parse_args()),indent=2))
    except Exception as e:print('Export:',e,file=sys.stderr);sys.exit(1)
