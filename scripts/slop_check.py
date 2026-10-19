#!/usr/bin/env python3
"""Deck-level review for common machine-made presentation patterns ("slop").

Reads a PPTX (python-pptx), a studio scene JSON (deck.scene.json) or a built HTML
deck (Playwright) and reports concrete, located signals: accent stripes on cards,
title underlines, decorative ordinals, one skeleton repeated across the deck,
English eyebrows in a CJK deck, a color identity reused for unrelated items,
near-duplicate pages, small live text, production/status labels on the canvas and
tool metadata in speaker notes. For HTML decks it also reads computed styles for the
decorations that make a deck look machine-made: tiled dot/grid backgrounds, glows and
blurs, gradient text, frosted glass, shadowed cards, letter-spaced caps labels,
separator chains, emoji, chips, giant faded numerals, stacked letters, a colored
headline phrase that labels nothing, and ▲ used for increases in Japanese decks.

These are review prompts with a suggested repair, not a taste score. Each warning
needs either a fix or a one-line reason why the pattern carries meaning here.
`--check` exits 2 only for errors (tool metadata in notes, production copy).
"""
from pathlib import Path
import argparse,colorsys,json,re,statistics,sys
from collections import Counter,defaultdict

W,H=1280,720
ROOT=Path(__file__).resolve().parents[1]
CJK=re.compile(r'[\u3040-\u30ff\u3400-\u9fff\uac00-\ud7af]')
ORDINAL=re.compile(r'^\s*(?:0?[1-9]|1\d)[.)]?\s*$')
EYEBROW=re.compile(r'^[A-Z0-9][A-Z0-9 /·.&:+\-–—]{2,}$')
STATUS=re.compile(r'(?:登壇用|発表用|社内用|作業用)?(?:ドラフト|下書き)|\b(?:draft|work in progress|WIP)\b',re.I)
TOOL_NOTES=re.compile(r'\[Export\]|HTML motion flattened|SVG/image elements remain pictures')

def hexrgb(c):
    if not c or not isinstance(c,str):return None
    c=c.strip().lstrip('#')
    if len(c)==3:c=''.join(x*2 for x in c)
    if not re.fullmatch(r'[0-9a-fA-F]{6}',c):return None
    return tuple(int(c[i:i+2],16)/255 for i in (0,2,4))
def hue_bucket(c):
    rgb=hexrgb(c)
    if not rgb:return None
    h,l,s=colorsys.rgb_to_hls(*rgb)
    if s<.3 or l<.12 or l>.92:return None
    return int(((h*360)+15)//30)%12
def saturated(c):return hue_bucket(c) is not None
def lab(c):
    rgb=hexrgb(c);lin=[v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4 for v in rgb]
    X=(.4124*lin[0]+.3576*lin[1]+.1805*lin[2])/.95047;Y=.2126*lin[0]+.7152*lin[1]+.0722*lin[2];Z=(.0193*lin[0]+.1192*lin[1]+.9505*lin[2])/1.08883
    f=lambda t:t**(1/3) if t>.008856 else 7.787*t+16/116
    return (116*f(Y)-16,500*(f(X)-f(Y)),200*(f(Y)-f(Z)))
def delta_e(a,b):return sum((x-y)**2 for x,y in zip(lab(a),lab(b)))**.5
def palette_colors(path):
    d=json.loads(Path(path).read_text(encoding='utf-8'));cols=list(d.get('tokens',{}).values())
    data=d.get('data',{});cols+=[c for k in('categorical','sequential','diverging') for c in data.get(k,[])]+[data.get('highlight'),data.get('context')]
    return [c for c in cols if hexrgb(c)]

# ---------- loaders: every slide becomes {id,title,notes,elements:[{kind,x,y,w,h,fill,color,text,size}]}
def load_scene(path):
    d=json.loads(Path(path).read_text(encoding='utf-8'));sx=W/d.get('width',W);sy=H/d.get('height',H);slides=[]
    for i,s in enumerate(d['slides']):
        els=[]
        for n in s.get('elements',[]):
            t=n.get('type');base={'x':n['x']*sx,'y':n['y']*sy,'w':n['w']*sx,'h':n['h']*sy}
            if t=='text':els.append({**base,'kind':'text','text':str(n.get('text','')),'size':n.get('size',16)*sx,'color':n.get('color')})
            elif t in('rect','ellipse'):
                if n.get('fill') not in(None,'none'):els.append({**base,'kind':'shape','fill':n.get('fill')})
            elif t=='line':els.append({**base,'kind':'line','fill':n.get('stroke')})
            elif t in('image','svg'):els.append({**base,'kind':'image'})
        slides.append({'id':s.get('id',f'slide-{i+1}'),'title':s.get('title',''),'notes':s.get('notes',''),'elements':els})
    return slides

def load_pptx(path):
    from pptx import Presentation
    from pptx.util import Emu
    p=Presentation(path);sx=W/p.slide_width;sy=H/p.slide_height;pt=lambda v:v*96/72*(W/(p.slide_width/914400*96))
    def color(fmt):
        try:
            if fmt.type is not None and fmt.fore_color and fmt.fore_color.type is not None:return '#'+str(fmt.fore_color.rgb)
        except Exception:return None
    def walk(shapes,out):
        for sh in shapes:
            if sh.shape_type==6:walk(sh.shapes,out);continue
            if sh.width is None:continue
            box={'x':sh.left*sx,'y':sh.top*sy,'w':sh.width*sx,'h':sh.height*sy}
            if sh.shape_type==13:out.append({**box,'kind':'image'});continue
            text=sh.text_frame.text.strip() if sh.has_text_frame else ''
            fill=None
            try:fill=color(sh.fill)
            except Exception:pass
            line=None
            try:line=('#'+str(sh.line.color.rgb)) if sh.line.fill.type==1 else None
            except Exception:pass
            if text:
                sizes=[r.font.size.pt for para in sh.text_frame.paragraphs for r in para.runs if r.font.size]
                rc=None
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        try:
                            if r.font.color and r.font.color.type is not None:rc='#'+str(r.font.color.rgb);break
                        except Exception:pass
                    if rc:break
                out.append({**box,'kind':'text','text':text,'size':pt(max(sizes)) if sizes else 24,'color':rc})
            if fill:out.append({**box,'kind':'line' if min(box['w'],box['h'])<1 else 'shape','fill':fill})
            elif line and (box['w']<2 or box['h']<2 or getattr(sh,'shape_type',None)==9):out.append({**box,'kind':'line','fill':line})
    slides=[]
    for i,s in enumerate(p.slides):
        els=[];walk(s.shapes,els)
        notes=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        title=next((e['text'] for e in sorted(els,key=lambda e:-e.get('size',0)) if e['kind']=='text'),'')
        slides.append({'id':f'slide-{i+1}','title':title.split('\n')[0],'notes':notes,'elements':els})
    return slides

HTML_JS=r'''()=>{const out=[];const slides=[...document.querySelectorAll('.slide')];
 const rgb=c=>{const m=c&&c.match(/rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?/);if(!m||(m[4]!==undefined&&+m[4]<.3))return null;return '#'+[m[1],m[2],m[3]].map(v=>(+v|0).toString(16).padStart(2,'0')).join('')};
 for(const [i,s] of slides.entries()){window.slideStudio?.show?.(i,999);window.slideStudio?.finish?.();const r0=s.getBoundingClientRect(),k=1280/r0.width,els=[];
  for(const el of s.querySelectorAll('*')){const cs=getComputedStyle(el);if(cs.visibility==='hidden'||cs.display==='none'||+cs.opacity===0)continue;const r=el.getBoundingClientRect();if(!r.width&&!r.height)continue;
   const box={x:(r.left-r0.left)*k,y:(r.top-r0.top)*k,w:r.width*k,h:r.height*k};
   const own=[...el.childNodes].filter(n=>n.nodeType===3).map(n=>n.textContent).join('').trim();
   if(own&&!(el instanceof SVGElement&&el.tagName!=='text'&&el.tagName!=='tspan'))els.push({...box,kind:'text',text:el.innerText?.trim()||own,size:parseFloat(cs.fontSize)*k,color:rgb(cs.color)});
   const bg=rgb(cs.backgroundColor);if(bg&&!(el===s))els.push({...box,kind:'shape',fill:bg});
   for(const side of ['Top','Right','Bottom','Left']){const bw=parseFloat(cs['border'+side+'Width']);const bc=rgb(cs['border'+side+'Color']);if(bw>=2&&bc&&cs['border'+side+'Style']!=='none'){const others=['Top','Right','Bottom','Left'].filter(x=>x!==side).every(x=>parseFloat(cs['border'+x+'Width'])<1);if(others){const t=bw*k;const b=side==='Top'?{x:box.x,y:box.y,w:box.w,h:t}:side==='Bottom'?{x:box.x,y:box.y+box.h-t,w:box.w,h:t}:side==='Left'?{x:box.x,y:box.y,w:t,h:box.h}:{x:box.x+box.w-t,y:box.y,w:t,h:box.h};els.push({...b,kind:'shape',fill:bc,edgeOf:{...box}})}}}
   if(el instanceof SVGGeometryElement){const f=rgb(cs.fill)||rgb(cs.stroke);if(f&&['rect','line'].includes(el.tagName))els.push({...box,kind:el.tagName==='line'?'line':'shape',fill:f,svg:true})}}
  out.push({id:s.id||('slide-'+(i+1)),title:s.getAttribute('aria-label')||'',notes:s.dataset.notes||'',elements:els,style:styleSignals(s,r0,k)})}return out;
 // Style signals of the machine-made look: ornament, glow, chips, spaced caps, emoji, an accent phrase that labels nothing.
 function styleSignals(s,r0,k){
  const rgba=c=>{const m=c&&c.match(/rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?/);return m?{hex:'#'+[m[1],m[2],m[3]].map(v=>(+v|0).toString(16).padStart(2,'0')).join(''),a:m[4]===undefined?1:+m[4]}:null};
  const label=el=>{const t=(el.textContent||'').trim().replace(/\s+/g,' ').slice(0,32);return el.tagName.toLowerCase()+(el.classList&&el.classList.length?'.'+[...el.classList].slice(0,2).join('.'):'')+(t?' "'+t+'"':'')};
  const st={pattern:[],glow:[],gradientText:[],glass:[],shadows:[],spacedCaps:[],dotChains:[],emoji:[],pills:[],faded:[],stacked:[],accent:null,marks:[],texts:[]};
  const area=r0.width*r0.height,h1=s.querySelector('h1');
  for(const el of [s,...s.querySelectorAll('*')]){
   const cs=getComputedStyle(el);if(cs.visibility==='hidden'||cs.display==='none'||+cs.opacity===0)continue;
   const r=el.getBoundingClientRect();if(!r.width||!r.height)continue;
   const own=[...el.childNodes].filter(n=>n.nodeType===3).map(n=>n.textContent).join('').trim(),fs=parseFloat(cs.fontSize)*k;
   const bi=cs.backgroundImage||'none';
   if(bi!=='none'&&/gradient\(/.test(bi)){
    const tiled=cs.backgroundSize.split(',').some(v=>{const q=v.trim().split(/\s+/).map(parseFloat);return q[0]>0&&q[0]*k<=48&&(q.length<2||isNaN(q[1])||q[1]*k<=48)});
    if(tiled||/repeating-/.test(bi))st.pattern.push(label(el));
    else if(/radial-gradient/.test(bi)&&r.width*r.height>=.15*area)st.glow.push(label(el)+' (radial wash)');}
   const f=cs.filter||'',bl=f.match(/blur\(([\d.]+)px\)/);
   if(bl&&+bl[1]>=6&&el.tagName!=='IMG')st.glow.push(label(el)+' (blur)');
   if(/drop-shadow/.test(f))st.glow.push(label(el)+' (drop-shadow)');
   if(el.tagName==='feGaussianBlur'&&+el.getAttribute('stdDeviation')>=4)st.glow.push('svg blur filter');
   if(cs.textShadow&&cs.textShadow!=='none'){const n=(cs.textShadow.match(/(-?[\d.]+)px/g)||[]).map(parseFloat);if(n.length>=3&&n[2]>=8)st.glow.push(label(el)+' (text glow)')}
   if(cs.webkitBackgroundClip==='text'||cs.backgroundClip==='text')st.gradientText.push(label(el));
   if(cs.backdropFilter&&cs.backdropFilter!=='none')st.glass.push(label(el));
   if(cs.boxShadow&&cs.boxShadow!=='none'){const n=(cs.boxShadow.match(/(-?[\d.]+)px/g)||[]).map(parseFloat);if(n.length>=3&&n[2]>=12)st.shadows.push(label(el))}
   if(own){
    const ls=parseFloat(cs.letterSpacing)||0,caps=cs.textTransform==='uppercase'||(/[A-Z]{3}/.test(own)&&own===own.toUpperCase());
    if(caps&&fs<=20&&ls/parseFloat(cs.fontSize)>=.08)st.spacedCaps.push(own.slice(0,40));
    if(fs<=20&&(own.match(/\s[·•|]\s/g)||[]).length>=2)st.dotChains.push(own.slice(0,60));
    if(/\p{Emoji_Presentation}|\uFE0F/u.test(own))st.emoji.push(own.slice(0,30));
    const rad=parseFloat(cs.borderTopLeftRadius)||0,boxed=(rgba(cs.backgroundColor)?.a||0)>.05||parseFloat(cs.borderTopWidth)>=1;
    if(boxed&&r.height*k<=44&&rad>=r.height/2-1&&own.length<=28)st.pills.push(own.slice(0,28));
    const col=rgba(cs.color);if(fs>=140&&((col&&col.a<=.35)||+cs.opacity<=.35))st.faded.push(own.slice(0,12));
    if(!(h1&&h1.contains(el))&&col)st.texts.push(col.hex);}
   if(el.children.length===0){const lines=(el.innerText||'').split('\n').map(x=>x.trim()).filter(Boolean);if(lines.length>=3&&lines.every(l=>l.length===1&&/[A-Za-z]/.test(l)))st.stacked.push(lines.join(''))}
   if(h1&&h1.contains(el))continue;
   if(el instanceof SVGGeometryElement){for(const p of [cs.fill,cs.stroke]){const c=rgba(p);if(c&&c.a>.3)st.marks.push(c.hex)}}
   else if(el!==s){const c=rgba(cs.backgroundColor);if(c&&c.a>.3)st.marks.push(c.hex);for(const sd of ['Top','Left','Bottom','Right'])if(parseFloat(cs['border'+sd+'Width'])>=1){const c2=rgba(cs['border'+sd+'Color']);if(c2&&c2.a>.3)st.marks.push(c2.hex)}}}
  if(h1){const hc=rgba(getComputedStyle(h1).color),kids=[...h1.querySelectorAll('*')].filter(e=>{const c=rgba(getComputedStyle(e).color);return c&&hc&&c.hex!==hc.hex&&(e.textContent||'').trim()});
   if(kids.length)st.accent={text:kids.map(e=>e.textContent.trim()).join(' ').slice(0,48),color:rgba(getComputedStyle(kids[0]).color).hex,head:hc&&hc.hex};}
  st.marks=[...new Set(st.marks)].slice(0,300);st.texts=[...new Set(st.texts)].slice(0,100);return st}}'''

def load_html(path):
    import glob
    from playwright.sync_api import sync_playwright
    exe=__import__('os').environ.get('SLIDE_STUDIO_CHROMIUM') or next(iter(glob.glob('/opt/pw-browsers/chromium-*/chrome-linux*/chrome')),None)
    with sync_playwright() as p:
        b=p.chromium.launch(**({'executable_path':exe} if exe else {}))
        try:
            pg=b.new_page(viewport={'width':1280,'height':800});pg.route('http*://**',lambda r:r.abort());pg.goto(Path(path).resolve().as_uri());pg.wait_for_timeout(300)
            slides=pg.evaluate(HTML_JS)
            try:
                meta=pg.evaluate('()=>{const m=document.getElementById("slide-studio-meta")||document.getElementById("slide-studio-scene");return m?JSON.parse(m.textContent):null}')
                if meta and isinstance(meta.get('slides'),list):
                    for s,m in zip(slides,meta['slides']):s['notes']=m.get('notes','') or s['notes'];s['title']=m.get('title') or s['title']
            except Exception:pass
            return slides
        finally:b.close()

def load(path):
    p=Path(path);suf=p.suffix.lower()
    if suf in('.pptx','.potx'):return load_pptx(p),'pptx'
    if suf=='.json':return load_scene(p),'scene'
    if suf in('.html','.htm'):return load_html(p),'html'
    raise ValueError('Input must be .pptx, .scene.json or .html')

# ---------- signals
def texts(s):return [e for e in s['elements'] if e['kind']=='text']
def fills(s):return [e for e in s['elements'] if e['kind'] in('shape','line')]
def is_card(e):return e['kind']=='shape' and e['w']>=120 and e['h']>=48 and e['w']*e['h']<0.7*W*H

def accent_stripes(s):
    found=[];cards=[e for e in s['elements'] if is_card(e)]
    for e in fills(s):
        thin,long=min(e['w'],e['h']),max(e['w'],e['h'])
        if thin>8 or long<16 or long<3*max(thin,1) or not saturated(e.get('fill')):continue
        host=e.get('edgeOf')
        if not host:
            for c in cards:
                if c is e:continue
                horiz=e['w']>=e['h']
                if horiz and abs(e['y']-c['y'])<=4 and e['x']>=c['x']-4 and e['x']+e['w']<=c['x']+c['w']+4 and e['w']>=.5*c['w']:host=c;break
                if not horiz and abs(e['x']-c['x'])<=4 and e['y']>=c['y']-4 and e['y']+e['h']<=c['y']+c['h']+4 and e['h']>=.5*c['h']:host=c;break
        if host:found.append(('card-edge',e));continue
        # a short rule under a heading, or a tick in front of a top-of-page label
        near=[t for t in texts(s) if (e['w']>=e['h'] and t['size']>=28 and 0<=e['y']-(t['y']+t['h'])<=40 and abs(e['x']-t['x'])<=24 and e['w']<=.6*W)
              or (e['h']>e['w'] and e['h']<=60 and e['y']<0.2*H and 0<=t['x']-(e['x']+e['w'])<=24 and abs(e['y']-t['y'])<=20)]
        if near:found.append(('beside-text',e))
    return found

def ordinals(s):
    ts=texts(s);body=statistics.median([t['size'] for t in ts]) if ts else 16
    return [t for t in ts if ORDINAL.match(t['text']) and t['size']>=max(24,1.4*body)]

def card_row(s):
    """Largest set of >=3 equal-width boxes (filled cards or text columns) side by side on one baseline."""
    boxes=[e for e in s['elements'] if is_card(e) or (e['kind']=='text' and 120<=e['w']<=0.34*W)];best=0
    for a in boxes:
        grp=[b for b in boxes if b['kind']==a['kind'] and abs(b['y']-a['y'])<=8 and abs(b['w']-a['w'])<=0.04*a['w']]
        xs=sorted(b['x'] for b in grp);cols=[x for k,x in enumerate(xs) if k==0 or x-xs[k-1]>=0.8*a['w']]
        if len(cols)>=3:best=max(best,len(cols))
    return best

def skeleton(s):
    ts=texts(s);top=[t for t in ts if t['y']<0.14*H and t['size']<=18]
    title=max((t for t in ts if t['y']<0.3*H),key=lambda t:t['size'],default=None)
    foot=any(t['y']>0.9*H and t['size']<=16 for t in ts)
    return (bool(top),round(title['y']/40) if title else None,card_row(s)>=3,foot)

def tokens(text):
    t=text.lower();words=set(re.findall(r'[a-z][a-z0-9\-]{2,}',t))
    cjk=''.join(CJK.findall(t));return words|{cjk[i:i+2] for i in range(len(cjk)-1)}

def review(slides,source,palette=None):
    findings=[];add=lambda sev,sig,slide,msg,fix:findings.append({'severity':sev,'signal':sig,'slide':slide,'message':msg,'repair':fix})
    heads=[h['text'] for h in (max((t for t in texts(s) if t['y']<0.3*H and not ORDINAL.match(t['text'])),key=lambda t:t['size'],default=None) for s in slides) if h]
    cjk_deck=bool(heads) and sum(bool(CJK.search(h)) for h in heads)>=0.5*len(heads)  # judged by page titles, not by product names
    rules=[]
    try:rules=[(r['id'],re.compile(r['pattern'],re.I if 'i' in r.get('flags','') else 0),r['level']) for r in json.loads((ROOT/'assets/copy-rules.json').read_text(encoding='utf-8'))['rules']]
    except Exception:pass
    sizes=[t['size'] for s in slides for t in texts(s)]
    for i,s in enumerate(slides):
        sid=s['id']
        for kind,e in accent_stripes(s):
            add('warning','accent-stripe',sid,f"Thin colored bar {round(e['w'])}×{round(e['h'])}px at ({round(e['x'])},{round(e['y'])}) {'on a card edge' if kind=='card-edge' else 'beside a label or under a title'}.",
                'Delete it. Separate cards with space or a tint; mark a real category with a direct label or the data mark itself.')
        o=ordinals(s)
        if len(o)>=2:add('warning','decorative-ordinals',sid,f"{len(o)} large ordinal numbers ({', '.join(t['text'] for t in o[:4])}).",
            'Keep numbers only for a real sequence the audience must follow; otherwise let the headings carry the order.')
        for t in texts(s):
            if cjk_deck and t['y']<0.14*H and t['size']<=18 and EYEBROW.match(t['text']) and sum(c.isalpha() for c in t['text'])>=4:
                add('warning','latin-eyebrow',sid,f"English eyebrow '{t['text'][:40]}' in a Japanese/CJK deck.",'Remove it, or use a label in the deck language that tells the reader something.')
            m=STATUS.search(t['text'])
            if m:add('error' if re.match(r'(?:登壇用|発表用|社内用|作業用)',m.group(0)) else 'warning','status-on-canvas',sid,f"Production/status word on the canvas: '{m.group(0)}' in '{t['text'][:40]}'.",'Move status to the file name, notes or runbook.')
            for rid,rx,lvl in rules:
                if rid=='ja-status-label':continue  # reported as status-on-canvas above
                mm=rx.search(t['text'])
                if mm:add('error' if lvl=='error' else 'warning',f'copy:{rid}',sid,f"'{mm.group(0)}' in '{t['text'][:50]}'.",'Rewrite for the audience; see references/writing-and-editing.md.')
        if TOOL_NOTES.search(s.get('notes') or ''):
            add('error','tool-metadata-in-notes',sid,'Speaker notes contain export/tool metadata.','Delete it; notes are spoken or read by the presenter.')
    # deck-level: one skeleton everywhere
    body=[s for s in slides[1:]] if len(slides)>3 else slides
    sig=Counter(skeleton(s) for s in body)
    if body and sig:
        (shape,n),=sig.most_common(1)
        if shape[0] and shape[2] and n>=4 and n>=0.6*len(body):add('warning','repeated-skeleton','deck',f"{n} of {len(body)} content slides share one layout skeleton (eyebrow={shape[0]}, card row={shape[2]}, footer={shape[3]}).",
            'Give each page the form its evidence needs: a chart, a table, a diagram, a single statement. Repeat a structure only for true comparisons.')
    rowslides=[s['id'] for s in body if card_row(s)>=3]
    if len(rowslides)>=3 and len(rowslides)>=0.5*len(body):add('warning','card-rows-everywhere','deck',f"Equal-width card rows on {len(rowslides)} of {len(body)} content slides: {', '.join(rowslides)}.",
        'Replace card rows with the relationship itself: a table for attributes, a chart for quantities, a flow for steps, prose for an argument.')
    # color identity reused for unrelated items
    hue_labels=defaultdict(lambda:defaultdict(set))
    named=lambda l:sum(c.isalpha() for c in l)>=3
    for s in slides:
        for t in texts(s):
            hb=hue_bucket(t.get('color'))
            if hb is not None and len(t['text'])<=28 and 0.1*H<t['y']<0.9*H:hue_labels[hb][t['text'].strip()].add(s['id'])
    for hb,labels in hue_labels.items():
        ident=[l for l,ss in labels.items() if len(ss)>=2 and named(l)]
        others=sorted({l for l in labels if l not in ident and not any(i in l or l in i for i in ident) and (named(l) or ORDINAL.match(l))})
        if ident and len(others)>=2:
            add('warning','color-identity-reused','deck',f"A hue used as the identity of '{ident[0]}' also colors unrelated items: {', '.join(others[:4])}.",
                'Keep identity colors for their entity only; use neutral ink or one accent for everything else.')
    # many hues on text: identity or decoration colors spread across the deck
    text_hues=defaultdict(set)
    for s in slides:
        for t in texts(s):
            hb=hue_bucket(t.get('color'))
            if hb is not None:text_hues[hb].add(s['id'])
    spread=[h for h,ss in text_hues.items() if len(ss)>=2]
    if len(spread)>=4:add('warning','rainbow-text','deck',f"Text is set in {len(spread)} different saturated hues across the deck.",
        'Keep text in ink and muted ink; give one accent to the single thing each page is about, and keep categorical colors on data marks.')
    if palette:
        allowed=palette_colors(palette);off=defaultdict(set)
        for s in slides:
            for e in s['elements']:
                c=e.get('color') if e['kind']=='text' else e.get('fill')
                if hexrgb(c) and min(delta_e(c,a) for a in allowed)>10:off[c.lower()].add(s['id'])
        for c,ss in sorted(off.items(),key=lambda kv:-len(kv[1]))[:8]:
            add('warning','off-palette',', '.join(sorted(ss)[:4]),f"{c} is not in the declared color system ({Path(palette).stem}).",'Map it to the nearest role (ink, inkMuted, rule, accent, a data color) or justify it.')
    # near-duplicate pages
    toks=[tokens(' '.join(t['text'] for t in texts(s))) for s in slides]
    for i in range(len(slides)):
        for j in range(i+1,len(slides)):
            a,b=toks[i],toks[j]
            if len(a)>=12 and len(b)>=12:
                jac=len(a&b)/len(a|b)
                if jac>=0.45:add('warning','near-duplicate-pages',f"{slides[i]['id']}+{slides[j]['id']}",f"The two pages share {round(jac*100)}% of their words.",'Merge them, or give the second page a different question to answer.')
    # machine-made style signals (HTML decks): aggregated per signal across the deck
    STYLE={
      'pattern':('ornamental-pattern','Tiled dot/grid/stripe background','Delete it. A slide background is paper: flat, quiet, carrying nothing the reader must decode.'),
      'glow':('glow-or-blur','Glow, blur or radial wash','Delete it. Emphasis comes from position, size, weight or the one accent color on the thing that matters.'),
      'gradientText':('gradient-text','Gradient-filled text','Set the text in ink or the single accent color.'),
      'glass':('glass-panel','Frosted-glass panel (backdrop-filter)','Remove the panel; put text on genuine negative space or a flat field.'),
      'spacedCaps':('spaced-caps-labels','Letter-spaced UPPERCASE labels','Use sentence case at the label size; keep caps for real abbreviations. Delete labels that only decorate.'),
      'dotChains':('separator-chain','Chain of "·" or "|" separators','Say less: keep the one fact the reader needs, or set the facts as a short sentence.'),
      'emoji':('emoji','Emoji on the canvas','Delete it; if an image is needed, draw a precise one that explains something.'),
      'pills':('chips','Pill-shaped chips/badges','Replace chips with plain words in the sentence or a table column; keep at most one status mark per page.'),
      'faded':('faded-numeral','Giant faded numeral used as decoration','Delete it; if the number matters, set it at full contrast and say what it counts.'),
      'stacked':('stacked-letters','Latin letters stacked vertically','Set the label horizontally or rotate the whole word; never stack letters.'),
    }
    for key,(sig,what,fix) in STYLE.items():
        hits=[(s['id'],x) for s in slides for x in (s.get('style') or {}).get(key,[])]
        if key=='spacedCaps':hits=[(sid,x) for sid,x in hits if sum(1 for t,_ in hits if t==sid)>=2]
        if key=='pills':hits=[(sid,x) for sid,x in hits if sum(1 for t,_ in hits if t==sid)>=3]
        if hits:
            ids=sorted({h[0] for h in hits},key=lambda x:[s['id'] for s in slides].index(x) if x in [s['id'] for s in slides] else 0)
            add('warning',sig,', '.join(ids[:6])+(' …' if len(ids)>6 else ''),f"{what} on {len(ids)} slide(s), e.g. {hits[0][1]}.",fix)
    shadow_slides=[s['id'] for s in slides if len((s.get('style') or {}).get('shadows',[]))>=3]
    if shadow_slides:add('warning','shadowed-cards',', '.join(shadow_slides[:6]),f"Soft drop shadows on 3+ boxes on {len(shadow_slides)} slide(s).",
        'Remove the shadows and most of the boxes; group with space and hairline rules. A box should represent an object.')
    for s in slides:
        a=(s.get('style') or {}).get('accent')
        if not a or not saturated(a['color']):continue
        anchored=[c for c in (s['style'].get('marks',[])+s['style'].get('texts',[])) if hexrgb(c) and delta_e(c,a['color'])<=12]
        if not anchored:add('warning','headline-accent-unanchored',s['id'],f"The colored headline phrase '{a['text']}' matches no mark or label on the page.",
            'Set the headline in one color, or color the phrase only when it names the same-colored mark it refers to.')
    tri=[]
    for s in slides:
        joined=' '.join(t['text'] for t in texts(s))
        m=re.search(r'▲\s*[+＋]\s*\d|[+＋]\s*\d[\d.,]*\s*%?\s*▲',joined)
        if m and (cjk_deck or CJK.search(joined)):tri.append((s['id'],m.group(0)))
    if True:
        if tri:add('warning','ja-triangle-for-increase',', '.join(sorted({x[0] for x in tri})[:6]),f"▲ used with a plus sign ('{tri[0][1][:30]}'). In Japanese business documents ▲ marks a negative number.",
            'Write increases as +1.2 and decreases as −1.2 (or ▲1.2 in an accounting table, explained once).')
    sizes=[t['size'] for s in slides for t in texts(s) if t['y']<0.9*H]
    if sizes:
        small=[x for x in sizes if x<14]
        if len(small)>=0.3*len(sizes):add('warning','small-live-text','deck',f"{len(small)} of {len(sizes)} text runs are under 14px on a 1280px canvas (about 10.5pt on a 13.33in slide).",
            'Cut duplicated words and enlarge what remains; move optional detail to notes or an appendix page.')
    errors=[f for f in findings if f['severity']=='error']
    return {'source':source,'slides':len(slides),'errors':len(errors),'warnings':len(findings)-len(errors),'findings':findings,
            'limits':['Geometry and text heuristics; they locate candidates, they do not judge taste or truth.','A warning is resolved by a fix or by a one-line reason the pattern carries meaning on that page.']}

def main():
    a=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter);a.add_argument('input');a.add_argument('--check',action='store_true',help='exit 2 when errors are found')
    a.add_argument('--palette',help='assets/color-systems/<id>.json the deck declares; colors outside it are reported')
    args=a.parse_args()
    try:
        slides,kind=load(args.input);r=review(slides,kind,args.palette)
        print(json.dumps(r,ensure_ascii=False,indent=2));return 2 if args.check and r['errors'] else 0
    except Exception as e:print('Slop check:',e,file=sys.stderr);return 1
if __name__=='__main__':sys.exit(main())
