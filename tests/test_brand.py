from pathlib import Path
import unittest,tempfile,json,sys,zipfile,importlib.util
from io import BytesIO
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from brand_common import validate_brand,local_logo

def module(name):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/(name+'.py'));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
create=module('brand-template').create;capture=module('brand-capture').capture
class BrandTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.brand=json.loads((ROOT/'assets/brands/example.json').read_text())
    def tearDown(self):self.tmp.cleanup()
    def test_native_six_layouts_and_editable_starters(self):
        p=self.root/'template.pptx';create(self.brand,p,self.root);d=Presentation(p);self.assertEqual(len(d.slide_layouts),6);self.assertEqual(len(d.slides),6);self.assertTrue(all(s.shapes.title.has_text_frame for s in d.slides));self.assertEqual(d.slide_layouts[2].name,'Paired Comparison')
    def test_potx_has_template_content_type(self):
        p=self.root/'template.potx';create(self.brand,p,self.root)
        with zipfile.ZipFile(p) as z:self.assertIn(b'presentationml.template.main+xml',z.read('[Content_Types].xml'))
    def test_theme_matches_tokens(self):
        p=self.root/'test.pptx';create(self.brand,p,self.root)
        with zipfile.ZipFile(p) as z:self.assertIn(self.brand['colors']['accent'][1:].encode(),z.read('ppt/theme/theme1.xml'))
    def test_inserted_slide_inherits_named_placeholders(self):
        p=self.root/'test.pptx';create(self.brand,p,self.root);d=Presentation(p);s=d.slides.add_slide(d.slide_layouts[2]);s.shapes.title.text='Actual comparison';self.assertEqual(len(s.placeholders),3);d.save(self.root/'reused.pptx');self.assertEqual(Presentation(self.root/'reused.pptx').slides[-1].shapes.title.text,'Actual comparison')
    def test_content_populates_without_template_captions(self):
        p=self.root/'test.pptx';create(self.brand,p,self.root,[{'layout':2,'title':'Observed difference','body':['A','B']}]);d=Presentation(p);text=' '.join(s.text for s in d.slides[0].shapes if s.has_text_frame);self.assertIn('Observed difference',text);self.assertNotIn('Replace with',text)
    def test_no_overwrite(self):
        p=self.root/'test.pptx';p.write_text('preserve');self.assertRaises(ValueError,create,self.brand,p,self.root);self.assertEqual(p.read_text(),'preserve')
    def test_capture_is_draft_and_excludes_private_text(self):
        p=self.root/'source.pptx';create(self.brand,p,self.root,[{'layout':0,'title':'PRIVATE-SENTINEL-TITLE','body':['PRIVATE-SENTINEL-CONTENT']}]);out=self.root/'private';capture(p,out);all_text=''.join(p.read_text() for p in out.iterdir());self.assertNotIn('PRIVATE-SENTINEL',all_text);b=json.loads((out/'brand.json').read_text());self.assertEqual(b['status'],'draft');self.assertRaises(ValueError,validate_brand,b)
    def test_capture_refuses_existing_directory(self):
        p=self.root/'test.pptx';create(self.brand,p,self.root);self.assertRaises(ValueError,capture,p,self.root)
    def test_logo_traversal_rejected(self):
        self.brand['logo']={'file':'../outside.png','alt':'logo','x':0,'y':0,'width':100,'height':50};self.assertRaises(ValueError,local_logo,self.brand,self.root)
    def test_placeholder_content_validated(self):
        for row in [{'layout':99,'title':'A'},{'layout':0,'title':''},{'layout':0,'title':'A','body':['A',5]}]:self.assertRaises(ValueError,create,self.brand,self.root/'bad.pptx',self.root,[row])
    def test_native_master_resets_centering_and_default_bullets(self):
        from lxml import etree
        p=self.root/'test.pptx';create(self.brand,p,self.root)
        ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        with zipfile.ZipFile(p) as z:
            xml=etree.fromstring(z.read('ppt/slideMasters/slideMaster1.xml'))
            levels=xml.findall('.//p:txStyles/p:titleStyle/a:lvl1pPr',ns)+xml.findall('.//p:txStyles/p:bodyStyle/a:lvl1pPr',ns)
            self.assertEqual(len(levels),2)
            for pp in levels:self.assertEqual(pp.get('algn'),'l');self.assertIsNotNone(pp.find('a:buNone',ns))
