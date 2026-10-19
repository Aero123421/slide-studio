import unittest,tempfile,json,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from slop_check import review,load_scene,load_pptx
def text(t,x,y,w,h,size,color='#1a1a1a'):return {'type':'text','text':t,'x':x,'y':y,'w':w,'h':h,'size':size,'color':color}
def rect(x,y,w,h,fill):return {'type':'rect','x':x,'y':y,'w':w,'h':h,'fill':fill}
HUES=['#8b5cf6','#14b8a6','#eab308','#f97316']
def sloppy_slide(i,ja=True):
    els=[rect(64,40,4,22,HUES[0]),text('0%d / HOW TO CHOOSE'%i,78,38,500,24,13),text('モデル選びは品質と時間と費用で決まる' if ja else 'Choosing a model balances quality and time',64,80,1100,60,40)]
    for k in range(4):
        x=64+k*290;els+= [rect(x,220,260,300,'#141b25'),rect(x,220,260,5,HUES[k]),text('0%d'%(k+1),x+20,240,80,60,44,HUES[k]),text(['GPT Alpha','GPT Beta','GPT Gamma','Model Delta'][k],x+20,320,220,30,20,HUES[k]),text(f'論点その{i}{k}',x+20,380,220,30,16,HUES[k])]
    els.append(text('FIELD GUIDE · 2026',64,690,400,20,12))
    return {'id':f's{i}','title':'t','layout':'freeform','notes':'Talk track.\n\n[Export] editable; HTML motion flattened.','elements':els}
def clean_slide(i):
    return {'id':f'c{i}','title':'t','layout':'freeform','notes':'Speak to the interval.','elements':[
        text(['受理1件あたりの費用はB案が最も低い','評価条件をそろえると順位が入れ替わる','次の四半期は2件の実務で比較する','価格は入力と出力で100倍の幅がある','結論：B案で小さく始める'][i],64,60,1100,60,40),
        rect(64+i*40,180,700,380,'#f3f4f6'),text('単位：USD / 100万トークン',64,600,500,24,18)]}
def scene(slides):return {'version':1,'width':1280,'height':720,'slides':slides}
def signals(r):return {f['signal'] for f in r['findings']}
class SlopCheck(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.dir=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def run_scene(self,slides):
        p=self.dir/'deck.scene.json';p.write_text(json.dumps(scene(slides),ensure_ascii=False));return review(load_scene(p),'scene')
    def test_generated_look_is_located(self):
        r=self.run_scene([sloppy_slide(i) for i in range(1,7)])
        for sig in ['accent-stripe','decorative-ordinals','latin-eyebrow','repeated-skeleton','card-rows-everywhere','tool-metadata-in-notes','color-identity-reused']:
            self.assertIn(sig,signals(r),sig)
        self.assertGreater(r['errors'],0)
    def test_varied_deck_without_patterns_is_quiet(self):
        r=self.run_scene([clean_slide(i) for i in range(5)])
        self.assertEqual(r['findings'],[])
    def test_status_label_on_canvas_is_an_error(self):
        s=clean_slide(0);s['elements'].append(text('登壇用ドラフト · 2026年9月23日時点',64,650,600,24,16))
        r=self.run_scene([s,clean_slide(1),clean_slide(2)])
        self.assertTrue(any(f['signal']=='status-on-canvas' and f['severity']=='error' for f in r['findings']))
    def test_data_bars_beside_small_labels_are_not_stripes(self):
        s=clean_slide(0)
        for k,w in enumerate([46,61,74]):s['elements']+= [text('段階%d'%k,560,280+k*88,120,24,15),rect(693,283+k*88,w,7,'#2563eb')]
        r=self.run_scene([s,clean_slide(1),clean_slide(2)])
        self.assertNotIn('accent-stripe',signals(r))
    def test_pptx_input_reads_shapes_colors_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches,Pt
        from pptx.dml.color import RGBColor
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        for i in range(5):
            s=prs.slides.add_slide(prs.slide_layouts[6])
            for k in range(3):
                card=s.shapes.add_shape(1,Inches(0.8+k*4),Inches(2.3),Inches(3.6),Inches(3));card.fill.solid();card.fill.fore_color.rgb=RGBColor(0x14,0x1b,0x25)
                bar=s.shapes.add_shape(1,Inches(0.8+k*4),Inches(2.3),Inches(3.6),Inches(0.05));bar.fill.solid();bar.fill.fore_color.rgb=RGBColor(0x8b,0x5c,0xf6)
            tb=s.shapes.add_textbox(Inches(0.8),Inches(0.8),Inches(11),Inches(1));tb.text_frame.text='各モデルの特徴をまとめる';tb.text_frame.paragraphs[0].runs[0].font.size=Pt(32)
            s.notes_slide.notes_text_frame.text='[Export] editable; HTML motion flattened.'
        p=self.dir/'deck.pptx';prs.save(p)
        r=review(load_pptx(p),'pptx')
        self.assertIn('accent-stripe',signals(r));self.assertIn('tool-metadata-in-notes',signals(r))
    def test_palette_discipline(self):
        r=self.run_scene([sloppy_slide(i) for i in range(1,5)])
        self.assertIn('rainbow-text',signals(r))
        s=clean_slide(0);s['elements'].append(rect(900,200,200,200,'#ff00aa'))
        r=review(load_scene(self._write([s,clean_slide(1),clean_slide(2)])),'scene',str(ROOT/'assets/color-systems/laboratory.json'))
        self.assertTrue(any(f['signal']=='off-palette' and '#ff00aa' in f['message'] for f in r['findings']))
        self.assertFalse(any(f['signal']=='off-palette' and '#f3f4f6' in f['message'] for f in r['findings']))
    def _write(self,slides):
        p=self.dir/'p.scene.json';p.write_text(json.dumps(scene(slides),ensure_ascii=False));return p
    def test_bundled_studio_examples_have_no_errors(self):
        import subprocess
        for name in ['technical','executive','talk']:
            out=self.dir/f'{name}.html'
            subprocess.run(['node',str(ROOT/'scripts/studio.mjs'),'build',str(ROOT/f'assets/examples/{name}.json'),'--out',str(out)],check=True,capture_output=True)
            r=review(load_scene(self.dir/f'{name}.scene.json'),'scene')
            self.assertEqual(r['errors'],0,name)
    def test_html_style_signals_of_the_machine_made_look(self):
        try:
            import playwright  # noqa: F401
            from slop_check import load_html
        except Exception:self.skipTest('Playwright is not installed')
        decorated=('<section class="slide" id="decorated" style="position:relative;width:1280px;height:720px;background:#f4f1ea radial-gradient(circle at 1px 1px,#bbb 1px,transparent 1.2px) 0 0/24px 24px">'
            '<h1 style="font-size:40px;color:#111">売上は計画を上回る一方、<em style="color:#d9442b;font-style:normal">新規ARRは未達</em></h1>'
            '<p style="font:600 13px Inter;letter-spacing:.14em;text-transform:uppercase">quarterly review</p><p style="font:600 13px Inter;letter-spacing:.14em;text-transform:uppercase">kpi summary</p>'
            + ''.join(f'<span style="display:inline-block;height:28px;line-height:28px;padding:0 12px;border-radius:14px;background:#e3f2ea">{t}</span>' for t in ['計画比 ▲ +2.2%','前年比','NEW'])
            + '<p style="font-size:20px">✨ 今期のハイライト</p></section>')
        anchored=('<section class="slide" id="anchored" style="position:relative;width:1280px;height:720px;background:#ffffff">'
            '<h1 style="font-size:40px;color:#111">ARR grew fastest in <em style="color:#2457c5;font-style:normal">the enterprise tier</em></h1>'
            '<svg width="600" height="200"><rect x="0" y="20" width="420" height="40" fill="#2457c5"/><rect x="0" y="80" width="300" height="40" fill="#9aa1ab"/></svg></section>')
        page=self.dir/'styled.html';page.write_text('<!doctype html><meta charset="utf-8"><body style="margin:0">'+decorated+anchored+'</body>',encoding='utf-8')
        try:slides=load_html(page)
        except Exception as e:self.skipTest(f'Chromium unavailable: {e}')
        r=review(slides,'html');sig=signals(r)
        for name in ['ornamental-pattern','spaced-caps-labels','chips','emoji','headline-accent-unanchored','ja-triangle-for-increase']:
            self.assertIn(name,sig,name)
        unanchored=[f['slide'] for f in r['findings'] if f['signal']=='headline-accent-unanchored']
        self.assertEqual(unanchored,['decorated'],'a colored phrase that names a same-colored mark is not flagged')
        lined=('<section class="slide" id="lined" style="position:relative;width:1280px;height:720px;background:#fff">'
            '<h1 style="font-size:40px;color:#111">Half the homes beat the <em style="color:#c2410c;font-style:normal">median of 2.9</em></h1>'
            '<svg width="600" height="240"><line x1="300" x2="300" y1="10" y2="220" stroke="#c2410c" stroke-width="2"/><circle cx="200" cy="120" r="5" fill="#9aa1ab"/></svg></section>'
            '<section class="slide" id="labelled" style="position:relative;width:1280px;height:720px;background:#fff">'
            '<h1 style="font-size:40px;color:#111">The <em style="color:#1d4ed8;font-style:normal">new model</em> catches more</h1>'
            '<svg width="600" height="240"><text x="20" y="40" font-size="20" fill="#1d4ed8">new model</text><path d="M0 200 L300 60" stroke="#9aa1ab" fill="none"/></svg></section>')
        page2=self.dir/'lines.html';page2.write_text('<!doctype html><meta charset="utf-8"><body style="margin:0">'+lined+'</body>',encoding='utf-8')
        r2=review(load_html(page2),'html')
        self.assertFalse([f for f in r2['findings'] if f['signal']=='headline-accent-unanchored'],'a same-coloured vertical line or SVG label anchors the phrase')
    def test_table_rules_are_typography_but_card_stripes_are_not(self):
        try:
            import playwright  # noqa: F401
            from slop_check import load_html
        except Exception:self.skipTest('Playwright is not installed')
        table=('<section class="slide" id="table" style="position:relative;width:1280px;height:720px;background:#fff">'
            '<h1 style="font-size:40px">Most teams need Growth</h1><table style="border-collapse:collapse;font-size:20px"><tr>'
            + ''.join(f'<th style="width:260px;text-align:left;border-top:3px solid {c};padding:10px 0">{t}</th>' for t,c in [('Starter','#fff'),('Growth','#0b6b4a'),('Scale','#fff')])
            + '</tr><tr><td>€49</td><td>€179</td><td>€490</td></tr></table></section>')
        cards=('<section class="slide" id="cards" style="position:relative;width:1280px;height:720px;background:#fff">'
            + ''.join(f'<div style="display:inline-block;width:300px;height:160px;margin:20px;background:#f3f5f8;border-left:4px solid #2457c5;border-radius:6px">{t}</div>' for t in ['Speed','Scale','Trust'])
            + '</section>')
        page=self.dir/'rules.html';page.write_text('<!doctype html><meta charset="utf-8"><body style="margin:0">'+table+cards+'</body>',encoding='utf-8')
        try:slides=load_html(page)
        except Exception as e:self.skipTest(f'Chromium unavailable: {e}')
        r=review(slides,'html')
        stripes={f['slide'] for f in r['findings'] if f['signal']=='accent-stripe'}
        self.assertNotIn('table',stripes,'a coloured rule over a table column is typography')
        self.assertIn('cards',stripes,'a coloured edge on tinted rounded cards is still flagged')
if __name__=='__main__':unittest.main()
